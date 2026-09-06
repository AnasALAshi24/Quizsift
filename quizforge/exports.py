from __future__ import annotations

import csv
from io import BytesIO, StringIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_BREAK

from .models import QuizQuestion, TermDefinition
from .questions import RANK_LABELS


LETTERS = "ABCD"


def paper_to_docx(questions: list[QuizQuestion], *, title: str, rank: int) -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(RANK_LABELS[rank])
    document.add_paragraph(f"Questions: {len(questions)}")

    for question in questions:
        document.add_heading(f"Question {question.number}", level=2)
        for prompt_line in question.prompt.splitlines():
            document.add_paragraph(prompt_line)
        for letter, choice in zip(LETTERS, question.choices):
            document.add_paragraph(f"{letter}. {choice}")

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Answer key", level=1)
    for question in questions:
        letter = LETTERS[question.choices.index(question.answer)]
        source = f" · source page {question.source_page}" if question.source_page else ""
        document.add_paragraph(f"{question.number}. {letter} — {question.answer}{source}")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def papers_to_zip(
    papers: dict[int, list[QuizQuestion]], *, title: str
) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for rank, questions in papers.items():
            archive.writestr(
                f"rank{rank}.docx",
                paper_to_docx(questions, title=title, rank=rank),
            )
    return output.getvalue()


def glossary_to_csv(entries: list[TermDefinition]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["term", "description", "source_page"])
    for entry in entries:
        writer.writerow([entry.term, entry.description, entry.source_page or ""])
    return output.getvalue().encode("utf-8-sig")
